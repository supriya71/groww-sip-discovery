"""Import REVIEW columns from review_workbook.xlsx into data/manual/corrections.csv.

Run from repo root: python -m pipeline.import_review
"""

from __future__ import annotations

import argparse
import csv
import sys
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

CORRECTIONS_PATH = REPO_ROOT / "data" / "manual" / "corrections.csv"
DEFAULT_WORKBOOK = REPO_ROOT / "data" / "review" / "review_workbook.xlsx"

from pipeline.corpus_filters import load_corrections  # noqa: E402
from pipeline.export import load_codebook  # noqa: E402
from pipeline.review_export import build_tagged_rows  # noqa: E402
from pipeline.review_workbook_common import SHEET_TAGGED, TAGGED_HEADERS  # noqa: E402
from pipeline.review_workbook_common import map_field_fix_to_correction, old_value_for_field  # noqa: E402

CORRECTION_FIELDS = ("id", "field", "old_value", "new_value", "note", "corrected_by", "date")


def _cell_map(row_cells: tuple, headers: tuple[str, ...]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for idx, name in enumerate(headers):
        if idx < len(row_cells):
            val = row_cells[idx]
            out[name] = val if val is None else str(val).strip()
        else:
            out[name] = ""
    return out


def _is_no(val: str) -> bool:
    return (val or "").strip().upper() in {"N", "NO"}


def load_existing_keys(rows: list[dict[str, str]]) -> set[tuple[str, str, str]]:
    return {(r["id"], r["field"], r["new_value"]) for r in rows}


def import_from_workbook(path: Path, *, corrected_by: str = "review_workbook") -> int:
    if not path.is_file():
        raise FileNotFoundError(f"Workbook not found: {path}")

    codebook = load_codebook()
    tagged_by_id = {r["id"]: r for r in build_tagged_rows(codebook, load_corrections())}

    wb = load_workbook(path, read_only=True, data_only=True)
    if SHEET_TAGGED not in wb.sheetnames:
        raise ValueError(f"Sheet {SHEET_TAGGED!r} missing in {path}")
    ws = wb[SHEET_TAGGED]

    existing = load_corrections()
    seen = load_existing_keys(existing)
    added: list[dict[str, str]] = []
    today = str(date.today())

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        cells = _cell_map(row, TAGGED_HEADERS)
        if not _is_no(cells.get("REVIEW: correct?", "")):
            continue
        field_label = cells.get("REVIEW: field to fix", "")
        new_raw = cells.get("REVIEW: correct value", "")
        note = cells.get("REVIEW: note", "") or "review_workbook correction"
        iid = cells.get("id", "")
        row_data = tagged_by_id.get(iid)
        if not row_data:
            print(f"  skip unknown id: {iid}")
            continue
        mapped = map_field_fix_to_correction(field_label, new_raw, row_data, codebook)
        if not mapped:
            print(f"  skip invalid correction row id={iid} field={field_label!r}")
            continue
        _iid, field, new_val = mapped
        old_val = old_value_for_field(
            row_data,
            "relevant" if field == "relevant" else field,
            codebook,
        )

        key = (_iid, field, new_val)
        if key in seen:
            continue
        seen.add(key)
        added.append(
            {
                "id": _iid,
                "field": field,
                "old_value": old_val,
                "new_value": new_val,
                "note": note,
                "corrected_by": corrected_by,
                "date": today,
            }
        )

    wb.close()
    if not added:
        print("No new corrections to add.")
        return 0

    CORRECTIONS_PATH.parent.mkdir(parents=True, exist_ok=True)
    write_header = not CORRECTIONS_PATH.is_file() or CORRECTIONS_PATH.stat().st_size == 0
    with CORRECTIONS_PATH.open("a", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CORRECTION_FIELDS)
        if write_header and not existing:
            writer.writeheader()
        for rec in added:
            writer.writerow(rec)

    print(f"Appended {len(added)} row(s) to {CORRECTIONS_PATH.relative_to(REPO_ROOT)}")
    return len(added)


def main() -> int:
    parser = argparse.ArgumentParser(description="Import REVIEW columns into corrections.csv")
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help="Path to review_workbook.xlsx",
    )
    parser.add_argument("--by", default="review_workbook", help="corrected_by value")
    args = parser.parse_args()
    try:
        import_from_workbook(args.workbook.resolve(), corrected_by=args.by)
    except (FileNotFoundError, ValueError) as exc:
        print(f"FAIL: {exc}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
