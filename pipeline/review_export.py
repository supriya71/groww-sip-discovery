"""Build human review Excel workbook under data/review/ (git-ignored).

Run from repo root: python -m pipeline.review_export
"""

from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

PROCESSED = REPO_ROOT / "data" / "processed"
REVIEW_DIR = REPO_ROOT / "data" / "review"
WORKBOOK_PATH = REVIEW_DIR / "review_workbook.xlsx"
CODEBOOK_PATH = REPO_ROOT / "codebook.yaml"

from pipeline.clean import iter_raw_jsonl_files  # noqa: E402
from pipeline.corpus_filters import apply_corrections, load_corrections  # noqa: E402
from pipeline.export import is_relevant, load_codebook, load_tag_rows, severity_score  # noqa: E402
from pipeline.filter import item_id, load_clean_rows  # noqa: E402
from pipeline.review_workbook_common import (  # noqa: E402
    DROPPED_HEADERS,
    REVIEW_HEADERS,
    SHEET_DROPPED,
    SHEET_SUMMARY,
    SHEET_TAGGED,
    SHEET_UNTAGGED,
    TAGGED_HEADERS,
    UNTAGGED_HEADERS,
    codebook_dropdown_values,
    FIELD_FIX_LABELS,
    row_link,
    tagged_row_values,
)
from pipeline.tag import load_filter_results, load_tag_finished_ids  # noqa: E402

HEADER_FILL = PatternFill("solid", fgColor="D9EAD3")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")
HEADER_FONT = Font(bold=True)


def load_clean_index() -> dict[str, dict[str, Any]]:
    return {item_id(r): r for r in load_clean_rows()}


def count_raw_pulled() -> Counter[str]:
    counts: Counter[str] = Counter()
    for path in iter_raw_jsonl_files():
        source_hint = path.parent.name
        for line in path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                row = json.loads(line)
                src = row.get("source") or source_hint
                if (row.get("text") or "").strip():
                    counts[src] += 1
    return counts


def funnel_by_source() -> dict[str, dict[str, int]]:
    pulled = count_raw_pulled()
    clean = Counter(r.get("source") or "unknown" for r in load_clean_rows())
    filt = load_filter_results()
    clean_idx = load_clean_index()
    filtered_by_src: Counter[str] = Counter()
    for iid, frec in filt.items():
        if frec.get("relevant") not in ("yes", "unsure", "no"):
            continue
        row = clean_idx.get(iid)
        if row:
            filtered_by_src[row.get("source") or "unknown"] += 1

    tagged_rows = build_tagged_rows(load_codebook(), load_corrections())
    tagged_by_src = Counter(r.get("source") or "unknown" for r in tagged_rows)
    relevant_by_src = Counter(
        r.get("source") or "unknown" for r in tagged_rows if is_relevant(r)
    )

    sources = sorted(
        set(pulled) | set(clean) | set(filtered_by_src) | set(tagged_by_src)
    )
    out: dict[str, dict[str, int]] = {}
    for src in sources:
        out[src] = {
            "pulled": pulled.get(src, 0),
            "cleaned": clean.get(src, 0),
            "quick_filtered": filtered_by_src.get(src, 0),
            "tagged": tagged_by_src.get(src, 0),
            "relevant": relevant_by_src.get(src, 0),
        }
    return out


def build_tagged_rows(codebook: dict[str, Any], corrections: list[dict[str, str]]) -> list[dict[str, Any]]:
    clean = load_clean_index()
    rows: list[dict[str, Any]] = []
    for rec in load_tag_rows():
        iid = rec["id"]
        base = clean.get(iid, {})
        merged = {
            **rec,
            "text": base.get("text") or rec.get("quote") or "",
            "context": base.get("context"),
            "year": base.get("year") or rec.get("year"),
            "source": base.get("source") or rec.get("source"),
            "url": base.get("url") or rec.get("url"),
            "app_id": base.get("app_id"),
            "video_id": base.get("video_id"),
        }
        merged = apply_corrections(merged, corrections)
        rows.append(merged)
    return rows


def build_untagged_rows() -> list[dict[str, Any]]:
    clean = load_clean_index()
    filt = load_filter_results()
    done = load_tag_finished_ids()
    rows: list[dict[str, Any]] = []
    for iid, frec in filt.items():
        if frec.get("relevant") not in ("yes", "unsure"):
            continue
        if iid in done:
            continue
        row = clean.get(iid)
        if not row:
            continue
        rows.append(
            {
                "id": iid,
                "source": row.get("source"),
                "year": row.get("year"),
                "link": row_link(row),
                "text": row.get("text") or "",
                "context": row.get("context") or "",
                "filter_relevant": frec.get("relevant"),
                "filter_reason": frec.get("reason") or "",
            }
        )
    rows.sort(key=lambda r: (r.get("source") or "", -(r.get("year") or 0)))
    return rows


def load_dropped_rows() -> list[dict[str, Any]]:
    path = PROCESSED / "dropped.jsonl"
    if not path.is_file():
        return []
    rows: list[dict[str, Any]] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        rows.append(row)
    rows.sort(key=lambda r: (r.get("source") or "", r.get("drop_reason") or ""))
    return rows


def write_sheet_header(ws, headers: tuple[str, ...]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL if not title.startswith("REVIEW:") else REVIEW_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def apply_table_style(ws, headers: tuple[str, ...], nrows: int) -> None:
    review_start = None
    for idx, h in enumerate(headers, start=1):
        if h.startswith("REVIEW:"):
            review_start = idx
            break
    wrap_cols = set()
    for name in ("post text", "context", "evidence quote", "detail"):
        if name in headers:
            wrap_cols.add(headers.index(name) + 1)
    for row in range(2, nrows + 2):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            if col in wrap_cols:
                cell.alignment = WRAP
            if review_start and col >= review_start:
                cell.fill = REVIEW_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2, nrows + 1)}"


def add_review_validations(ws, codebook: dict[str, Any], nrows: int) -> None:
    if nrows < 1:
        nrows = 5000
    last_row = max(nrows + 1, 5000)
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    yn.error = "Enter Y or N"
    ws.add_data_validation(yn)
    col_correct = TAGGED_HEADERS.index("REVIEW: correct?") + 1
    col_field = TAGGED_HEADERS.index("REVIEW: field to fix") + 1
    col_value = TAGGED_HEADERS.index("REVIEW: correct value") + 1
    letter_c = get_column_letter(col_correct)
    letter_f = get_column_letter(col_field)
    letter_v = get_column_letter(col_value)
    yn.add(f"{letter_c}2:{letter_c}{last_row}")

    codes = codebook_dropdown_values(codebook)
    hidden = ws.parent.create_sheet("_codebook_lists")
    hidden.sheet_state = "hidden"
    for i, val in enumerate(codes, start=1):
        hidden.cell(row=i, column=1, value=val)
    for i, val in enumerate(FIELD_FIX_LABELS, start=1):
        hidden.cell(row=i, column=2, value=val)

    dv_field = DataValidation(
        type="list",
        formula1=f"='_codebook_lists'!$B$1:$B${len(FIELD_FIX_LABELS)}",
        allow_blank=True,
    )
    ws.add_data_validation(dv_field)
    dv_field.add(f"{letter_f}2:{letter_f}{last_row}")

    dv_val = DataValidation(
        type="list",
        formula1=f"='_codebook_lists'!$A$1:$A${len(codes)}",
        allow_blank=True,
    )
    ws.add_data_validation(dv_val)
    dv_val.add(f"{letter_v}2:{letter_v}{last_row}")


def write_summary(ws, funnel: dict[str, dict[str, int]]) -> None:
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    instructions = [
        "SIP discovery — human review workbook",
        "",
        "How to use:",
        "1. Review rows on 'All tagged posts' (sorted by source, severity, year).",
        "2. Fill REVIEW columns when a tag is wrong: correct?=N, field, correct value, note.",
        "3. Run: python -m pipeline.import_review",
        "4. Re-export dashboard: python -m pipeline.export",
        "",
        "Counts below are per source at each pipeline stage.",
    ]
    for i, line in enumerate(instructions, start=1):
        ws.cell(row=i, column=1, value=line)

    start = len(instructions) + 2
    headers = ["source", "pulled", "cleaned", "quick-filtered", "tagged", "relevant"]
    stage_keys = ["pulled", "cleaned", "quick_filtered", "tagged", "relevant"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=start, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL

    row = start + 1
    totals = {k: 0 for k in stage_keys}
    for src in sorted(funnel):
        stages = funnel[src]
        ws.cell(row=row, column=1, value=src)
        for col, key in enumerate(stage_keys, start=2):
            val = stages[key]
            ws.cell(row=row, column=col, value=val)
            totals[key] += val
        row += 1

    ws.cell(row=row, column=1, value="TOTAL").font = HEADER_FONT
    for col, key in enumerate(stage_keys, start=2):
        ws.cell(row=row, column=col, value=totals[key]).font = HEADER_FONT


def set_column_widths(ws, headers: tuple[str, ...]) -> None:
    widths = {
        "id": 36,
        "source": 14,
        "year": 8,
        "link": 40,
        "post text": 55,
        "context": 28,
        "evidence quote": 40,
        "REVIEW: note": 32,
    }
    for idx, name in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 16)


def run_export() -> Path:
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    codebook = load_codebook()
    corrections = load_corrections()
    tagged = build_tagged_rows(codebook, corrections)
    tagged.sort(
        key=lambda r: (
            r.get("source") or "",
            -severity_score(r),
            -(r.get("year") or 0),
        )
    )

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = SHEET_SUMMARY
    write_summary(ws_sum, funnel_by_source())

    ws_tag = wb.create_sheet(SHEET_TAGGED)
    write_sheet_header(ws_tag, TAGGED_HEADERS)
    for i, row in enumerate(tagged, start=2):
        for col, val in enumerate(tagged_row_values(row, codebook), start=1):
            ws_tag.cell(row=i, column=col, value=val)
    apply_table_style(ws_tag, TAGGED_HEADERS, len(tagged))
    set_column_widths(ws_tag, TAGGED_HEADERS)
    add_review_validations(ws_tag, codebook, len(tagged))

    ws_un = wb.create_sheet(SHEET_UNTAGGED)
    write_sheet_header(ws_un, UNTAGGED_HEADERS)
    untagged = build_untagged_rows()
    for i, rec in enumerate(untagged, start=2):
        vals = [
            rec["id"],
            rec["source"],
            rec["year"],
            rec["link"],
            rec["text"],
            rec["context"],
            rec["filter_relevant"],
            rec["filter_reason"],
        ]
        for col, val in enumerate(vals, start=1):
            ws_un.cell(row=i, column=col, value=val)
    apply_table_style(ws_un, UNTAGGED_HEADERS, len(untagged))
    set_column_widths(ws_un, UNTAGGED_HEADERS)

    ws_drop = wb.create_sheet(SHEET_DROPPED)
    write_sheet_header(ws_drop, DROPPED_HEADERS)
    dropped = load_dropped_rows()
    for i, row in enumerate(dropped, start=2):
        vals = [
            item_id(row),
            row.get("source"),
            row.get("year"),
            row_link(row),
            row.get("text") or "",
            row.get("drop_reason") or "",
            row.get("detail") or "",
        ]
        for col, val in enumerate(vals, start=1):
            ws_drop.cell(row=i, column=col, value=val)
    apply_table_style(ws_drop, DROPPED_HEADERS, len(dropped))
    set_column_widths(ws_drop, DROPPED_HEADERS)

    wb.save(WORKBOOK_PATH)
    return WORKBOOK_PATH


def main() -> int:
    path = run_export()
    print(f"Wrote {path.relative_to(REPO_ROOT)}")
    print(f"  tagged rows: {len(build_tagged_rows(load_codebook(), load_corrections()))}")
    print(f"  not yet tagged: {len(build_untagged_rows())}")
    print(f"  dropped: {len(load_dropped_rows())}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
