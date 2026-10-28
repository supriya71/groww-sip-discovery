"""Score filled labelling sheet vs AI tags; write agreement JSON + evidence log.

Run after filling data/review/labelling_sheet.xlsx:
  python -m pipeline.labelling_score
Then: python -m pipeline.export
"""

from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

WORKBOOK_PATH = REPO_ROOT / "data" / "review" / "labelling_sheet.xlsx"
MANIFEST_PATH = REPO_ROOT / "data" / "manual" / "labelling_sample.json"
AGREEMENT_PATH = REPO_ROOT / "data" / "manual" / "labelling_agreement.json"

from pipeline.export import load_codebook, load_tag_rows, merge_records  # noqa: E402
from pipeline.filter import item_id, load_clean_rows  # noqa: E402
from pipeline.labelling_sheet_common import (  # noqa: E402
    FIELD_LABELS,
    HUMAN_FIELDS,
    LABELLING_HEADERS,
    SHEET_LABELLING,
    ai_values,
    append_labelling_evidence_log,
    display_ai,
    display_human,
    enrich_from_clean,
    fields_agree,
    human_values,
)


def load_ai_index() -> dict[str, dict[str, Any]]:
    clean = {item_id(r): r for r in load_clean_rows()}
    return {
        r["id"]: enrich_from_clean(r, clean)
        for r in merge_records(load_tag_rows(), clean)
    }


def row_filled(cells: dict[str, str]) -> bool:
    return any((cells.get(f) or "").strip() for f in HUMAN_FIELDS)


def score_workbook(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"Labelling workbook missing: {path}")

    codebook = load_codebook()
    ai_index = load_ai_index()
    manifest = {}
    if MANIFEST_PATH.is_file():
        manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))

    wb = load_workbook(path, read_only=True, data_only=True)
    if SHEET_LABELLING not in wb.sheetnames:
        raise ValueError(f"Sheet {SHEET_LABELLING!r} not found")
    ws = wb[SHEET_LABELLING]

    field_stats = {f: {"agree": 0, "total": 0} for f in HUMAN_FIELDS}
    disagreements: list[dict[str, Any]] = []
    scored_ids: list[str] = []

    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or not row[0]:
            continue
        cells = {
            LABELLING_HEADERS[i]: (row[i] if i < len(row) and row[i] is not None else "")
            for i in range(len(LABELLING_HEADERS))
        }
        if isinstance(cells.get("id"), str):
            cells["id"] = cells["id"].strip()
        else:
            cells["id"] = str(cells.get("id") or "").strip()
        if not cells["id"]:
            continue
        if not row_filled(cells):
            continue

        ai_row = ai_index.get(cells["id"])
        if not ai_row:
            continue
        human = human_values({k: str(v) for k, v in cells.items()}, codebook)
        ai = ai_values(ai_row, codebook)
        scored_ids.append(cells["id"])

        for field in HUMAN_FIELDS:
            if not human.get(field) and field != "stated_triggers":
                continue
            if field == "stated_triggers" and not human.get("stated_triggers"):
                continue
            field_stats[field]["total"] += 1
            if fields_agree(field, human[field], ai[field]):
                field_stats[field]["agree"] += 1
            else:
                disagreements.append(
                    {
                        "id": cells["id"],
                        "field": FIELD_LABELS.get(field, field),
                        "field_key": field,
                        "human": display_human(codebook, field, human[field]),
                        "ai": display_ai(codebook, field, ai[field]),
                        "snippet": (cells.get("post text") or "")[:120],
                        "stratum": next(
                            (
                                it.get("stratum")
                                for it in manifest.get("items") or []
                                if it.get("id") == cells["id"]
                            ),
                            None,
                        ),
                    }
                )

    wb.close()

    fields_out: dict[str, Any] = {}
    total_agree = 0
    total_n = 0
    for field, st in field_stats.items():
        n = st["total"]
        a = st["agree"]
        total_agree += a
        total_n += n
        fields_out[field] = {
            "label": FIELD_LABELS.get(field, field),
            "agree": a,
            "total": n,
            "pct": round(100.0 * a / n, 1) if n else None,
        }

    report = {
        "scored_at": datetime.now(timezone.utc).isoformat(),
        "workbook": str(path.relative_to(REPO_ROOT)).replace("\\", "/"),
        "sample_size": len(scored_ids),
        "expected_sample_size": manifest.get("sample_size"),
        "strata_note": manifest.get("strata_note", ""),
        "overall_pct": round(100.0 * total_agree / total_n, 1) if total_n else None,
        "fields": fields_out,
        "disagreements": disagreements,
        "scored_ids": scored_ids,
    }
    return report


def main() -> int:
    try:
        report = score_workbook(WORKBOOK_PATH.resolve())
    except (FileNotFoundError, ValueError) as exc:
        print(f"FAIL: {exc}")
        return 1

    if report["sample_size"] == 0:
        print("No filled human labels found (fill columns D–I on Labelling sheet, row 3+).")
        return 1

    AGREEMENT_PATH.parent.mkdir(parents=True, exist_ok=True)
    AGREEMENT_PATH.write_text(json.dumps(report, indent=2), encoding="utf-8")
    append_labelling_evidence_log(report)

    print(f"Scored {report['sample_size']} items — overall {report['overall_pct']}% field agreement")
    print(f"  Wrote {AGREEMENT_PATH.relative_to(REPO_ROOT)}")
    print(f"  Disagreements: {len(report['disagreements'])}")
    print("  Run: python -m pipeline.export  (refresh Validation page)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
